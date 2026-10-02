import re
from docx import Document
from docx.shared import Pt, RGBColor

doc = Document()

style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(12)

def add_formatted_paragraph(doc, text):
    p = doc.add_paragraph()
    pattern = r'(\*\*\*.*?\*\*\*|\*\*.*?\*\*|\*.*?\*|[^*]+)'
    tokens = re.findall(pattern, text)
    for token in tokens:
        if token.startswith('***') and token.endswith('***'):
            run = p.add_run(token[3:-3])
            run.bold = True
            run.italic = True
            run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
        elif token.startswith('**') and token.endswith('**'):
            run = p.add_run(token[2:-2])
            run.bold = True
        elif token.startswith('*') and token.endswith('*'):
            run = p.add_run(token[1:-1])
            run.italic = True
        else:
            p.add_run(token)
    return p

doc.add_heading('EP03 - Texto para el Narrador', 0)
doc.add_heading('El ataque de la IA: El plan de China para sobrevivir', 2)
doc.add_paragraph()

idx = doc.add_paragraph()
r1 = idx.add_run('Indice de formato: ')
r1.bold = True
r2 = idx.add_run('Negrita')
r2.bold = True
idx.add_run(' = enfasis, pausar despues  |  ')
r3 = idx.add_run('Cursiva')
r3.italic = True
idx.add_run(' = ironia/chiste  |  ')
r4 = idx.add_run('Negrita+cursiva en rojo')
r4.bold = True
r4.italic = True
r4.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
idx.add_run(' = cambio de emocion base')

doc.add_paragraph('-' * 60)

lines = [
    "Si lees las noticias hoy, parece que el **ataque final** de la Inteligencia Artificial es inminente. Ingenieros de Estados Unidos renuncian en masa, aterrorizados de que nazca **Skynet**, rogando que los gobiernos intervengan. ***asombro*** Y frente a este apocalipsis, China tiene un plan muy distinto para sobrevivir.",
    "Pero spoiler: no estan construyendo bunkers contra robots. Mientras el mundo occidental cae en el panico, hoy te voy a demostrar como el gigante asiatico esta usando este teatro para ganar **la guerra del siglo veintiuno**.",
    "-" * 60,
    "Hoy te voy a demostrar que el **capitalismo** mas poderoso del mundo esta rogando por regulacion estatal como si fuera socialismo, mientras el **comunismo** mas famoso del planeta acelera sin restricciones morales como el capitalista mas salvaje de la historia. *Si, leiste bien.* Dejame tu me gusta, suscribite y decime en los comentarios: quien crees que gana esta carrera?",
    "-" * 60,
    "Vamos a los datos. En los ultimos meses, renunciaron ingenieros clave en Anthropic y en OpenAI. Todos con el mismo mensaje: la Inteligencia Artificial nos puede **exterminar** en menos de diez anos. Piden a gritos que los gobiernos regulen, que intervengan, que frenen el desarrollo antes de que sea demasiado tarde.",
    "Ahora pausemos un segundo y preguntemonos algo que nadie se esta preguntando: ***indignacion*** a quien le conviene exactamente este panico?",
    "Si convences al gobierno de que la Inteligencia Artificial es tan peligrosa como una **central nuclear**, los requisitos para desarrollarla se vuelven inalcanzables. Licencias, auditorias, comites de etica, estandares de seguridad imposibles de pagar. Y adivina que pasa: las startups emergentes y el codigo abierto se mueren ahogados en burocracia. Solo sobreviven los tres o cuatro gigantes que ya tienen abogados, lobbyistas y capital para cumplir cualquier regulacion que se invente.",
    "Es exactamente la misma jugada que hicieron con la industria del **cannabis legal**. *Resulta que la marihuana medicinal es un negocio serio que necesita estandares farmaceuticos.* Exigieron regulaciones tan costosas que los pequenos productores se fundieron, y las megacorporaciones se quedaron con todo el mercado. El miedo moral se convierte en su mejor **modelo de negocios**.",
    "Occidente no esta jugando a salvar a la humanidad. Esta jugando al **monopolio**.",
    "-" * 60,
    "Y ahora crucemos el oceano. Escuchaste a algun ingeniero de DeepSeek renunciar por conflictos morales? ***asombro*** No. Ni uno. Y aca esta la paradoja que te rompe la cabeza.",
    "Mucha gente cree que China no regula nada. **Falso.** China tiene uno de los marcos regulatorios de Inteligencia Artificial mas prescriptivos del mundo. Desde dos mil veinticinco, todo contenido generado por IA debe estar etiquetado obligatoriamente: texto, audio, imagen, video. Los deepfakes y la clonacion de voz tienen controles tecnicos estrictos.",
    "Pero y aca esta **la clave** China regula **el resultado**, no el desarrollo. Controlan lo que la maquina le dice a su gente. Jamas le pondrian un limite de velocidad al **codigo** en si mismo.",
    "Para ellos, la Inteligencia Artificial no es un dilema etico de ciencia ficcion. Es una **carrera armamentista**. El Plan Quinquenal dos mil veintiseis a dos mil treinta tiene a la IA en el centro de la estrategia industrial del pais. El Estado chino no es un arbitro moral intentando frenar el progreso: es el **socio principal** que te dice: *controla el cerebro a los ciudadanos, pero en el laboratorio, pisa el acelerador a fondo.*",
    "Eso es lo que el capitalismo liberal occidental llama, *ironicamente*, falta de etica.",
    "-" * 60,
    "Mientras Occidente usa el miedo al apocalipsis para construir un muro regulatorio que protege a **tres monopolios corporativos**, Oriente avanza sin frenos morales para quedarse con el control tecnologico del proximo siglo.",
    "No es que China no tenga miedo. Es que entienden algo que sus competidores occidentales se niegan a admitir: en una carrera armamentista, **el que se detiene a filosofar, pierde**. Y el que gana esa carrera no va a hacerlo disparando misiles. Lo va a hacer controlando la infraestructura tecnologica que va a decidir como funciona el mundo en los proximos **cincuenta anos**.",
    "***indignacion*** El verdadero riesgo de la Inteligencia Artificial no es que una maquina tome conciencia y nos dispare rayos laser. El verdadero peligro es que nos quedemos debatiendo el apocalipsis mientras cedemos todo nuestro **futuro tecnologico**.",
    "-" * 60,
    "Nos mienten con el miedo, nos venden la regulacion como salvacion, y mientras miramos la pantalla esperando a Terminator, el tablero ya se esta moviendo. No nos va a hundir la tecnologia. Nos va a hundir nuestra propia **hipocresia**.",
    "Si queres seguir entendiendo las jugadas que los grandes medios no te explican, suscribite al canal y activa la campanita. Y contame en los comentarios: crees que Occidente puede recuperar terreno en esta carrera, o ya es demasiado tarde?",
]

for line in lines:
    if line.startswith('-'):
        doc.add_paragraph(line)
    else:
        add_formatted_paragraph(doc, line)

doc.save(r'01_guiones\EP03_IA_CHINA\EP03_TEXTO_NARRADOR.docx')
print('DOCX creado OK')
